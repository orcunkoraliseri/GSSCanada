"""
T14 — WP12.1 follow-up: activity codes observed in the diaries vs codebook rows.

Task doc: 2J_docs_occ_nTemp/writing/submission/rejection revision/impl/2026-09-15_T14_wp12_codes_observed.md

Run on Speed only (see task doc). Expects to be run with cwd =
/speed-scratch/o_iseri/2J_revision/T14/, reading:
    input/episode_{2005,2010,2015,2022}.csv
    input/Data Harmonization_activityCategories - execution.xlsx
and writing:
    T14_out/t14_code_counts.csv   (cycle, measure, value, detail)

No installs. If openpyxl is not importable, writes T14_out/NOT_FOUND.txt and stops.
"""

import os
import sys

IN_DIR = "input"
OUT_DIR = "T14_out"

ACTIVITY_EXCEL = os.path.join(IN_DIR, "Data Harmonization_activityCategories - execution.xlsx")
ACT_SHEET_MAP = {2005: "2005codebook", 2010: "2010codebook", 2015: "2015codebook", 2022: "2022codebook"}
EPISODE_FILES = {y: os.path.join(IN_DIR, f"episode_{y}.csv") for y in ACT_SHEET_MAP}
SENTINEL_CODE = {2005: 995, 2010: 995, 2015: 95, 2022: 9999}  # 02_harmonizeGSS.py:420-425

os.makedirs(OUT_DIR, exist_ok=True)

try:
    import openpyxl  # noqa: F401
except ImportError:
    with open(os.path.join(OUT_DIR, "NOT_FOUND.txt"), "w") as f:
        f.write("openpyxl not importable in job env (T14 task doc step: stop, no installs)\n")
    print("NOT FOUND: openpyxl not importable; stopping per task doc.")
    sys.exit(0)

import pandas as pd


def build_activity_crosswalks(excel_path):
    """Verbatim copy of build_activity_crosswalks, 02_harmonizeGSS.py:372-391.
    Extended (only) to also return the raw kept-row count per cycle for R1 —
    the original function only returns `crosswalks`."""
    wb = openpyxl.load_workbook(excel_path, data_only=True)
    crosswalks = {}
    row_counts = {}
    for cycle_year, sheet_name in ACT_SHEET_MAP.items():
        if sheet_name not in wb.sheetnames:
            continue
        ws = wb[sheet_name]
        lookup = {}
        rows_kept = 0
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row[0] is None or row[1] is None or type(row[0]) == str:
                continue
            category = int(row[0])
            raw_str = str(row[1]).replace(",", ".")
            raw_code = float(raw_str) if cycle_year == 2010 else int(float(raw_str))
            lookup[raw_code] = category
            if cycle_year == 2010:
                lookup[int(raw_code)] = category
            rows_kept += 1
        crosswalks[cycle_year] = lookup
        row_counts[cycle_year] = rows_kept
    return crosswalks, row_counts


records = []


def add(cycle, measure, value, detail=""):
    records.append({"cycle": cycle, "measure": measure, "value": value, "detail": detail})
    print(f"[{cycle}] {measure} = {value}  {detail}")


def write_out():
    out_path = os.path.join(OUT_DIR, "t14_code_counts.csv")
    pd.DataFrame(records, columns=["cycle", "measure", "value", "detail"]).to_csv(out_path, index=False)
    print(f"WROTE {out_path}")


def main():
    if not os.path.isfile(ACTIVITY_EXCEL):
        add("ALL", "R0_input_missing", "NOT FOUND", ACTIVITY_EXCEL)
        write_out()
        return

    crosswalks, row_counts = build_activity_crosswalks(ACTIVITY_EXCEL)

    for cycle in (2005, 2010, 2015, 2022):
        lookup = crosswalks.get(cycle)
        if lookup is None:
            add(cycle, "R0_sheet_missing", "NOT FOUND")
            continue

        # raw column per 02_harmonizeGSS.py:398
        raw_col = "ACTCODE" if cycle in (2005, 2010) else "TUI_01"
        ep_path = EPISODE_FILES[cycle]
        if not os.path.isfile(ep_path):
            add(cycle, "R0_episode_file_missing", "NOT FOUND", ep_path)
            continue

        # R1 — codebook rows read (rows kept by the build_activity_crosswalks filter)
        add(cycle, "R1_codebook_rows_read", row_counts[cycle])

        # R2 — distinct keys in the lookup; 2010 split float vs int alias keys
        if cycle == 2010:
            float_keys = sum(1 for k in lookup if type(k) is float)
            int_keys = sum(1 for k in lookup if type(k) is int)
            add(cycle, "R2_distinct_keys_float", float_keys)
            add(cycle, "R2_distinct_keys_int", int_keys)
            add(cycle, "R2_distinct_keys_total", len(lookup))
        else:
            add(cycle, "R2_distinct_keys", len(lookup))

        df = pd.read_csv(ep_path, usecols=[raw_col])

        if cycle == 2010:
            # 02_harmonizeGSS.py:404-416 (float primary, int fallback)
            parsed = pd.to_numeric(df[raw_col], errors="coerce").astype(float)
            matched = parsed.map(lookup)
            mask = matched.isna() & df[raw_col].notna()
            if mask.any():
                fallback = (
                    pd.to_numeric(df.loc[mask, raw_col], errors="coerce")
                    .fillna(-1)
                    .astype(int)
                    .map(lookup)
                )
                matched.loc[mask] = fallback
        else:
            # 02_harmonizeGSS.py:418
            parsed = pd.to_numeric(df[raw_col], errors="coerce")
            matched = parsed.map(lookup)

        observed_mask = parsed.notna()
        observed = parsed[observed_mask]
        matched_obs = matched[observed_mask]

        # R3 — distinct non-null raw codes observed in episodes
        distinct_codes = observed.unique()
        add(cycle, "R3_distinct_observed_codes", len(distinct_codes))

        code_df = pd.DataFrame({"code": observed.values, "matched": matched_obs.notna().values})
        grp = code_df.groupby("code").agg(episodes=("code", "size"), any_matched=("matched", "any"))

        # R4 — observed codes that ARE in the codebook
        r4_codes = grp.index[grp["any_matched"]]
        add(cycle, "R4_observed_in_codebook", len(r4_codes))

        # R6 — observed codes NOT in the codebook, with episode counts
        r6 = grp.loc[~grp["any_matched"]]
        r6_list = ";".join(f"{float(c):g}:{int(n)}" for c, n in r6["episodes"].items())
        add(cycle, "R6_observed_not_in_codebook", len(r6), r6_list)

        # R5 — codebook codes never observed (list them)
        codebook_codes_norm = sorted(set(float(k) for k in lookup.keys()))
        observed_codes_norm = set(float(c) for c in distinct_codes)
        never_observed = [c for c in codebook_codes_norm if c not in observed_codes_norm]
        add(cycle, "R5_codebook_never_observed", len(never_observed),
            ";".join(f"{c:g}" for c in never_observed))

        # R7 — does code 2 occur in this cycle's episodes, with count
        count2 = int((observed == 2.0).sum())
        add(cycle, "R7_code2_occurs", count2 > 0, f"count={count2}")

        # note special codes 995/95/9999 handled at 02_harmonizeGSS.py:420-425
        sc = SENTINEL_CODE[cycle]
        sc_count = int((observed == sc).sum())
        add(cycle, "note_sentinel_code_override", sc, f"episodes_with_this_code={sc_count}")

    write_out()


if __name__ == "__main__":
    main()
