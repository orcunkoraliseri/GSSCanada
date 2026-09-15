import openpyxl

PATH = r"C:\Users\o_iseri\Desktop\GSSCanada\GSSCanada-main\2J_docs_occ_nTemp\references_activityCodes\Data Harmonization_activityCategories - execution.xlsx"

SHEET_MAP = {
    2005: "2005codebook",
    2010: "2010codebook",
    2015: "2015codebook",
    2022: "2022codebook",
}

wb = openpyxl.load_workbook(PATH, data_only=True)
print("sheets in workbook:", wb.sheetnames)

for cycle, sheet_name in SHEET_MAP.items():
    ws = wb[sheet_name]
    all_rows = list(ws.iter_rows(min_row=2, values_only=True))
    total_rows_with_maxrow = ws.max_row - 1  # rows from row2..max_row

    # Raw: any row where col B (index1) is not None (regardless of col A)
    raw_nonempty_b = [r for r in all_rows if r[1] is not None]

    # Pipeline filter exactly as build_activity_crosswalks:
    # skip if row[0] is None or row[1] is None or type(row[0]) == str
    kept = []
    skipped_reason = {"colA_none": 0, "colB_none": 0, "colA_str": 0}
    for r in all_rows:
        if r[0] is None:
            skipped_reason["colA_none"] += 1
            continue
        if r[1] is None:
            skipped_reason["colB_none"] += 1
            continue
        if type(r[0]) == str:
            skipped_reason["colA_str"] += 1
            continue
        kept.append(r)

    # Build the lookup dict exactly as pipeline does, to get final unique-key count
    lookup = {}
    dup_overwrites = []
    for r in kept:
        category = int(r[0])
        raw_str = str(r[1]).replace(",", ".")
        raw_code = float(raw_str) if cycle == 2010 else int(float(raw_str))
        if raw_code in lookup and lookup[raw_code] != category:
            dup_overwrites.append((raw_code, lookup[raw_code], category))
        elif raw_code in lookup:
            dup_overwrites.append((raw_code, lookup[raw_code], category, "same-cat-dup"))
        lookup[raw_code] = category
        if cycle == 2010:
            lookup[int(raw_code)] = category

    print(f"\n=== cycle {cycle} (sheet {sheet_name}) ===")
    print("ws.max_row:", ws.max_row, "-> data rows (max_row-1):", total_rows_with_maxrow)
    print("rows with non-empty col B (raw code present):", len(raw_nonempty_b))
    print("rows kept by pipeline filter (pre-dedup count):", len(kept))
    print("skip reasons:", skipped_reason)
    print("final unique lookup keys (what pipeline actually uses):", len(lookup))
    print("duplicate raw_code rows within sheet:", len(dup_overwrites))
    if dup_overwrites:
        print("  dup details (raw_code, prev_cat, new_cat):", dup_overwrites[:20])

    # list any rows skipped due to colA being string (to see what they are)
    str_colA_rows = [r for r in all_rows if r[0] is not None and type(r[0]) == str]
    if str_colA_rows:
        print("rows with col A as string (first 10):", str_colA_rows[:10])

    # list rows with colB present but colA None (orphan raw codes)
    orphan = [r for r in all_rows if r[0] is None and r[1] is not None]
    if orphan:
        print("rows with col A None but col B present (first 10):", orphan[:10])

    # rows with colA present but colB None
    orphanB = [r for r in all_rows if r[0] is not None and r[1] is None and type(r[0]) != str]
    if orphanB:
        print("rows with col A present, col B None (first 10):", orphanB[:10])
